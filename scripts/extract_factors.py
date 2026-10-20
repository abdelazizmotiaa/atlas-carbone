#!/usr/bin/env python3
"""
Atlas Carbone - Pipeline d'extraction et normalisation des facteurs d'émissions
Produit data/catalog_seed.json
"""

import os
import re
import json
import openpyxl

EXCEL_PATH = '/home/abdelaziz/Downloads/Bilan Carbone Maroc-compatible postes iso14064 et GHG Protocol-draft-15 octobre 2024 (3) (1).xlsx'
OUTPUT_JSON = '/home/abdelaziz/Desktop/PhD/analysis/atlas-carbone/data/catalog_seed.json'

SHEET_CONFIG = {
    'P1-S1-Source Fixe combustion': {
        'category': 'combustion_fixe',
        'scope': 1,
        'iso_category': 1,
        'ghg_category': 'Direct Scope 1'
    },
    'P2-S1-Source mobile Combust.': {
        'category': 'combustion_mobile',
        'scope': 1,
        'iso_category': 1,
        'ghg_category': 'Direct Scope 1'
    },
    'P3-S1-Emiss. Proc. hors énergie': {
        'category': 'procedes_industriels',
        'scope': 1,
        'iso_category': 1,
        'ghg_category': 'Direct Scope 1'
    },
    'P4-S1-Emiss. fugitives Climatis': {
        'category': 'emissions_fugitives',
        'scope': 1,
        'iso_category': 1,
        'ghg_category': 'Direct Scope 1'
    },
    'P5-S1-UTCATF': {
        'category': 'biomasse_sols_forets',
        'scope': 1,
        'iso_category': 1,
        'ghg_category': 'Direct Scope 1'
    },
    'P6-S2-ConsommationElectricité': {
        'category': 'electricite_reseau',
        'scope': 2,
        'iso_category': 2,
        'ghg_category': 'Indirect Scope 2'
    },
    'P7-S2-Emiss. Vap. Chal. Froid': {
        'category': 'vapeur_chaleur_froid',
        'scope': 2,
        'iso_category': 2,
        'ghg_category': 'Indirect Scope 2'
    },
    'P8-S3-En. non incluse post-1&7': {
        'category': 'energie_amont',
        'scope': 3,
        'iso_category': 4,
        'ghg_category': 'Scope 3 Cat 3: Énergie amont'
    },
    'P9-S3-.Achats prod. et services': {
        'category': 'achats_biens_services',
        'scope': 3,
        'iso_category': 4,
        'ghg_category': 'Scope 3 Cat 1: Biens et services achetés'
    },
    'P10-S3.Immobilisations de biens': {
        'category': 'immobilisations_biens',
        'scope': 3,
        'iso_category': 4,
        'ghg_category': 'Scope 3 Cat 2: Biens d équipement'
    },
    'P11-S3- Déchets': {
        'category': 'dechets_exploitations',
        'scope': 3,
        'iso_category': 4,
        'ghg_category': 'Scope 3 Cat 5: Déchets générés'
    },
    'P12-S3-Trans.March.amont': {
        'category': 'fret_marchandises_amont',
        'scope': 3,
        'iso_category': 3,
        'ghg_category': 'Scope 3 Cat 4: Fret amont'
    },
    'P13-S3-Transp. Professionnels': {
        'category': 'deplacements_professionnels',
        'scope': 3,
        'iso_category': 3,
        'ghg_category': 'Scope 3 Cat 6: Déplacements professionnels'
    },
    'P16-S3-Trans.Client.Visit': {
        'category': 'deplacements_visiteurs_clients',
        'scope': 3,
        'iso_category': 3,
        'ghg_category': 'Scope 3 Cat 9: Transport aval'
    },
    'P17-S3-Transport march.aval': {
        'category': 'fret_marchandises_aval',
        'scope': 3,
        'iso_category': 3,
        'ghg_category': 'Scope 3 Cat 9: Transport aval'
    },
    'P22-S3-Trajet dom.Travail.Empl.': {
        'category': 'deplacements_domicile_travail',
        'scope': 3,
        'iso_category': 3,
        'ghg_category': 'Scope 3 Cat 7: Déplacements domicile-travail'
    }
}

def clean_unit(u_str):
    if not u_str:
        return 'unite'
    u = u_str.strip()
    u = u.replace('Litre', 'litre').replace('Tonne', 'tonne').replace(' ', '')
    u = u.replace('kgCO2e/', '').replace('kgCO2/', '').replace('kgCO2e/(', '').replace(')', '')
    u = u.replace('tepPCI', 'tep').replace('kWhPCI', 'kWh')
    u = u.replace('passager.km', 'passager-km').replace('tonne.km', 'tonne-km')
    return u

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    all_factors = []
    factor_counter = 1

    seen_signatures = set()

    for sname, cfg in SHEET_CONFIG.items():
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        sheet_factors = []

        # Find rows where Col 2 is a text label, Col 3 is a positive float/int, Col 4 is an identifiable unit string
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 2).value
            val = ws.cell(r, 3).value
            unit_raw = ws.cell(r, 4).value
            uncertainty_raw = ws.cell(r, 5).value
            source_raw = ws.cell(r, 6).value if ws.max_column >= 6 else None

            # Skip header lines or empty labels
            if not label or not isinstance(label, str) or len(label.strip()) < 2:
                continue
            if not isinstance(val, (int, float)) or val <= 0:
                continue
            if not unit_raw or not isinstance(unit_raw, str) or not any(x in unit_raw.lower() for x in ['kg', 't', 'kwh', 'l', 'm', 'km', 'appareil']):
                continue

            clean_lbl = label.strip()
            # Determine perimeter and verification status
            is_amont_combustion = 'amont et combustion' in clean_lbl.lower()
            perimeter = 'combustion_et_amont' if is_amont_combustion else 'combustion'
            
            # Review status according to rules:
            # If in Scope 1 and marked "amont et combustion", status must be "to_verify" because lifecycle footprint shouldn't be 100% in Scope 1
            if cfg['scope'] == 1 and is_amont_combustion:
                status = 'to_verify'
                verification_note = 'Attention: Facteur mixte amont+combustion affecté au Scope 1. Doit être ventilé avant clôture définitive.'
            else:
                status = 'validated'
                verification_note = 'Facteur validé issu du référentiel national / Base Carbone.'

            # Handle uncertainty
            unc_pct = 0.20 # default 20%
            if isinstance(uncertainty_raw, (int, float)):
                unc_pct = float(uncertainty_raw)
            elif isinstance(uncertainty_raw, str):
                cleaned_unc = uncertainty_raw.replace('%', '').strip().replace(',', '.')
                try:
                    unc_pct = float(cleaned_unc) / 100.0 if '%' in uncertainty_raw else float(cleaned_unc)
                except ValueError:
                    unc_pct = 0.20

            unit_norm = clean_unit(unit_raw)
            factor_id = f"FE-MA-{factor_counter:04d}"
            factor_counter += 1

            sig = f"{clean_lbl.lower()}|{unit_norm}|{cfg['category']}"
            if sig in seen_signatures:
                # avoid pure duplicate factors
                continue
            seen_signatures.add(sig)

            factor_obj = {
                'id': factor_id,
                'version': '1.0',
                'name': clean_lbl,
                'category': cfg['category'],
                'scope': cfg['scope'],
                'iso_category': cfg['iso_category'],
                'ghg_category': cfg['ghg_category'],
                'value': float(val),
                'numerator_unit': 'kgCO2e',
                'denominator_unit': unit_norm,
                'raw_unit': str(unit_raw).strip(),
                'perimeter': perimeter,
                'geography': 'MA' if any(k in str(source_raw or '').lower() for k in ['fm6', 'maroc', 'onee']) else 'Monde',
                'reference_year': 2020 if '2020' in clean_lbl else 2023,
                'source': str(source_raw).strip() if source_raw else 'Base Carbone Maroc / FM6E',
                'sheet_origin': sname,
                'row_origin': r,
                'uncertainty_pct': unc_pct,
                'status': status,
                'verification_note': verification_note,
                'access_level': 'common_catalog'
            }
            sheet_factors.append(factor_obj)

        all_factors.extend(sheet_factors)
        print(f"[{sname}] Extraits : {len(sheet_factors)} facteurs")

    # Add official grid electricity location-based and market-based options explicitly
    all_factors.append({
        'id': f"FE-MA-{factor_counter:04d}",
        'version': '1.0',
        'name': "Électricité réseau national Maroc - Mix moyen (Location-based)",
        'category': 'electricite_reseau',
        'scope': 2,
        'iso_category': 2,
        'ghg_category': 'Indirect Scope 2',
        'value': 0.642602,
        'numerator_unit': 'kgCO2e',
        'denominator_unit': 'kWh',
        'raw_unit': 'kgCO2e/kWh',
        'perimeter': 'production_electrique',
        'geography': 'MA',
        'reference_year': 2023,
        'source': 'ONEE / Base Carbone Maroc (hors pertes)',
        'sheet_origin': 'P6-S2-ConsommationElectricité',
        'row_origin': 49,
        'uncertainty_pct': 0.05,
        'status': 'validated',
        'verification_note': 'Facteur officiel net hors pertes de distribution.',
        'access_level': 'common_catalog'
    })
    factor_counter += 1

    all_factors.append({
        'id': f"FE-MA-{factor_counter:04d}",
        'version': '1.0',
        'name': "Électricité réseau national Maroc - Avec pertes réseau 12.2% (Location-based)",
        'category': 'electricite_reseau',
        'scope': 2,
        'iso_category': 2,
        'ghg_category': 'Indirect Scope 2',
        'value': 0.721,
        'numerator_unit': 'kgCO2e',
        'denominator_unit': 'kWh',
        'raw_unit': 'kgCO2e/kWh',
        'perimeter': 'production_et_pertes',
        'geography': 'MA',
        'reference_year': 2020,
        'source': 'ONEE / FM6E (avec 12,2% pertes de transport)',
        'sheet_origin': 'P6-S2-ConsommationElectricité',
        'row_origin': 49,
        'uncertainty_pct': 0.10,
        'status': 'validated',
        'verification_note': 'Facteur historique brut pertes comprises.',
        'access_level': 'common_catalog'
    })
    factor_counter += 1

    all_factors.append({
        'id': f"FE-MA-{factor_counter:04d}",
        'version': '1.0',
        'name': "Électricité renouvelable certifiée (PPA / Autoconsommation) (Market-based)",
        'category': 'electricite_reseau',
        'scope': 2,
        'iso_category': 2,
        'ghg_category': 'Indirect Scope 2',
        'value': 0.0,
        'numerator_unit': 'kgCO2e',
        'denominator_unit': 'kWh',
        'raw_unit': 'kgCO2e/kWh',
        'perimeter': 'market_based_zero',
        'geography': 'MA',
        'reference_year': 2026,
        'source': 'GHG Protocol Market-based Guidance (avec certificat d origine garanti)',
        'sheet_origin': 'System',
        'row_origin': 0,
        'uncertainty_pct': 0.0,
        'status': 'validated',
        'verification_note': 'Applicable uniquement sur justification contractuelle (Certificat d Origine, PPA vert).',
        'access_level': 'common_catalog'
    })

    print(f"\nTOTAL FACTEURS EXTRAITS ET NORMALISÉS : {len(all_factors)}")
    with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(all_factors, f, indent=2, ensure_ascii=False)
    print(f"Fichier de seed généré : {OUTPUT_JSON}")

if __name__ == '__main__':
    main()
